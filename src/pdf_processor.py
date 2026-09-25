"""Extract slide-level content from PDFs and PowerPoint files locally."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
from typing import Any

from .config import PipelineConfig
from .pdf_runtime import muted_mupdf_errors
from .utils import clean_text, dump_json


SUPPORTED_SLIDE_SUFFIXES = {".pdf", ".ppt", ".pptx"}


class SlideProcessingError(RuntimeError):
    """Raised when a PDF or presentation cannot be turned into slide records."""


class PDFProcessor:
    """Create a consistent slide schema from PDF pages or PPTX slides.

    PDFs normally do not contain speaker notes.  For PPTX input the processor
    also reads PowerPoint speaker notes when they are present.
    """

    def __init__(self, config: PipelineConfig) -> None:
        self.config = config

    def process(self, presentation_path: str | Path, output_path: str | Path) -> dict[str, Any]:
        """Extract slide number, title, text, notes, and local visual OCR hints."""
        presentation_path = Path(presentation_path)
        output_path = Path(output_path)
        if not presentation_path.is_file():
            raise SlideProcessingError(f"Presentation file not found: {presentation_path}")
        suffix = presentation_path.suffix.lower()
        if suffix not in SUPPORTED_SLIDE_SUFFIXES:
            raise SlideProcessingError("Use a PDF, PPTX, or legacy PPT presentation file.")
        if suffix == ".pdf":
            slides = self._process_pdf(presentation_path, output_path.parent / "slide_images")
        elif suffix == ".pptx":
            slides = self._process_pptx(presentation_path, output_path.parent / "slide_images")
        else:
            converted_path = self._convert_legacy_ppt(presentation_path, output_path.parent / "converted")
            slides = self._process_pdf(converted_path, output_path.parent / "slide_images")

        payload = {
            "metadata": {
                "source_file": presentation_path.name,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "slide_count": len(slides),
                "format": suffix.lstrip("."),
            },
            "slides": slides,
        }
        dump_json(output_path, payload)
        return payload

    @staticmethod
    def _convert_legacy_ppt(path: Path, output_dir: Path) -> Path:
        """Convert a legacy PPT locally through LibreOffice when available."""
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            raise SlideProcessingError(
                "Legacy .ppt requires local LibreOffice for conversion. Install LibreOffice, "
                "or save the file as PDF/PPTX before uploading."
            )
        output_dir.mkdir(parents=True, exist_ok=True)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(path)],
                check=True,
                capture_output=True,
                text=True,
                timeout=180,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            raise SlideProcessingError(f"LibreOffice could not convert the PPT: {exc}") from exc
        converted = output_dir / f"{path.stem}.pdf"
        if not converted.is_file():
            raise SlideProcessingError("LibreOffice finished but did not create a PDF for the PPT file.")
        return converted

    def _process_pdf(self, path: Path, image_dir: Path) -> list[dict[str, Any]]:
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:
            raise SlideProcessingError("PyMuPDF is not installed. Run: pip install -r requirements.txt") from exc

        image_dir.mkdir(parents=True, exist_ok=True)
        slides: list[dict[str, Any]] = []
        try:
            with muted_mupdf_errors(fitz), fitz.open(path) as document:
                for number, page in enumerate(document, start=1):
                    raw_text = page.get_text("text")
                    lines = [clean_text(line) for line in raw_text.splitlines() if clean_text(line)]
                    title = self._infer_title(lines, number)
                    content = clean_text("\n".join(lines))
                    preview_path = image_dir / f"slide_{number:03d}_preview.png"
                    rendered = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
                    preview_path.write_bytes(rendered.tobytes("png"))
                    images, visual_text = self._extract_pdf_images(page, number, image_dir)
                    # Scanned slides often have no extractable page text; local OCR is a useful fallback.
                    if not content and self.config.enable_ocr:
                        ocr_render = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                        visual_text.append(self._ocr_bytes(ocr_render.tobytes("png")))
                        visual_text = [item for item in visual_text if item]
                        content = " ".join(visual_text)
                        title = self._infer_title(content.splitlines(), number)
                    slides.append(
                        self._slide_record(
                            number,
                            title,
                            content,
                            notes="",
                            images=images,
                            visual_text=visual_text,
                            preview_image=str(preview_path),
                        )
                    )
        except Exception as exc:
            raise SlideProcessingError(f"Could not process PDF: {exc}") from exc
        return slides

    def _process_pptx(self, path: Path, image_dir: Path) -> list[dict[str, Any]]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:
            raise SlideProcessingError("python-pptx is not installed. Run: pip install -r requirements.txt") from exc

        image_dir.mkdir(parents=True, exist_ok=True)
        rendered_previews = self._render_powerpoint_previews(path, image_dir)
        try:
            deck = Presentation(path)
        except Exception as exc:
            raise SlideProcessingError(f"Could not open PPTX: {exc}") from exc

        slides: list[dict[str, Any]] = []
        for number, slide in enumerate(deck.slides, start=1):
            title_shape = slide.shapes.title
            title = clean_text(title_shape.text) if title_shape and getattr(title_shape, "has_text_frame", False) else ""
            text_parts: list[str] = []
            images: list[str] = []
            visual_text: list[str] = []
            for image_index, shape in enumerate(self._walk_shapes(slide.shapes), start=1):
                if getattr(shape, "has_text_frame", False):
                    text = clean_text(shape.text)
                    if text and text != title:
                        text_parts.append(text)
                if getattr(shape, "has_table", False):
                    table_text = " | ".join(
                        clean_text(cell.text)
                        for row in shape.table.rows
                        for cell in row.cells
                        if clean_text(cell.text)
                    )
                    if table_text:
                        text_parts.append(f"Table: {table_text}")
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        extension = shape.image.ext or "png"
                        image_path = image_dir / f"slide_{number:03d}_image_{image_index}.{extension}"
                        image_path.write_bytes(shape.image.blob)
                        images.append(str(image_path))
                        hint = self._ocr_bytes(shape.image.blob)
                        if hint:
                            visual_text.append(hint)
                    except Exception:
                        # An unreadable embedded image should not prevent text extraction.
                        continue
            notes = ""
            try:
                notes_frame = slide.notes_slide.notes_text_frame
                notes = clean_text(notes_frame.text) if notes_frame else ""
            except Exception:
                notes = ""
            if not title:
                title = self._infer_title(text_parts, number)
            content = "\n".join(text_parts)
            preview_path = rendered_previews.get(number, image_dir / f"slide_{number:03d}_preview.png")
            if not preview_path.is_file():
                self._make_text_preview(preview_path, number, title, content)
            slides.append(
                self._slide_record(
                    number,
                    title,
                    content,
                    notes,
                    images,
                    visual_text,
                    preview_image=str(preview_path),
                )
            )
        return slides

    @staticmethod
    def _render_powerpoint_previews(path: Path, image_dir: Path) -> dict[int, Path]:
        """Render real PPTX slides through local LibreOffice, with a safe fallback."""
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            return {}
        rendered_dir = image_dir.parent / "rendered_presentation"
        rendered_dir.mkdir(parents=True, exist_ok=True)
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(rendered_dir), str(path)],
                check=True,
                capture_output=True,
                text=True,
                timeout=180,
            )
            import fitz

            output: dict[int, Path] = {}
            with muted_mupdf_errors(fitz), fitz.open(rendered_dir / f"{path.stem}.pdf") as document:
                for number, page in enumerate(document, start=1):
                    preview = image_dir / f"slide_{number:03d}_preview.png"
                    preview.write_bytes(
                        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).tobytes("png")
                    )
                    output[number] = preview
            return output
        except Exception:
            return {}

    @classmethod
    def _walk_shapes(cls, shapes: Any) -> Any:
        """Yield nested PowerPoint shapes so grouped text/images are not lost."""
        for shape in shapes:
            yield shape
            child_shapes = getattr(shape, "shapes", None)
            if child_shapes is not None:
                yield from cls._walk_shapes(child_shapes)

    def _extract_pdf_images(self, page: Any, number: int, image_dir: Path) -> tuple[list[str], list[str]]:
        images: list[str] = []
        visual_text: list[str] = []
        seen_xrefs: set[int] = set()
        document = page.parent
        for image_index, image_info in enumerate(page.get_images(full=True), start=1):
            xref = image_info[0]
            if xref in seen_xrefs:
                continue
            seen_xrefs.add(xref)
            try:
                extracted = document.extract_image(xref)
                extension = extracted.get("ext", "png")
                image_path = image_dir / f"slide_{number:03d}_image_{image_index}.{extension}"
                image_path.write_bytes(extracted["image"])
                images.append(str(image_path))
                hint = self._ocr_bytes(extracted["image"])
                if hint:
                    visual_text.append(hint)
            except Exception:
                continue
        return images, visual_text

    def _ocr_bytes(self, image_data: bytes) -> str:
        """Best-effort local OCR; absence of the Tesseract binary is harmless."""
        if not self.config.enable_ocr:
            return ""
        try:
            from PIL import Image
            import pytesseract

            return clean_text(pytesseract.image_to_string(Image.open(BytesIO(image_data))))
        except Exception:
            return ""

    @staticmethod
    def _make_text_preview(path: Path, number: int, title: str, content: str) -> None:
        """Create a readable local fallback when PowerPoint rendering is unavailable."""
        try:
            from PIL import Image, ImageDraw, ImageFont
            import textwrap

            image = Image.new("RGB", (1280, 720), "#f8fafc")
            draw = ImageDraw.Draw(image)
            title_font = ImageFont.truetype("Arial.ttf", 46)
            body_font = ImageFont.truetype("Arial.ttf", 28)
            small_font = ImageFont.truetype("Arial.ttf", 20)
        except (ImportError, OSError):
            from PIL import Image, ImageDraw, ImageFont
            import textwrap

            image = Image.new("RGB", (1280, 720), "white")
            draw = ImageDraw.Draw(image)
            title_font = body_font = small_font = ImageFont.load_default()
        draw.rectangle((0, 0, 1280, 14), fill="#2563eb")
        draw.text((72, 55), title or f"Slide {number}", fill="#0f172a", font=title_font)
        y = 145
        for paragraph in clean_text(content).split(" • "):
            for line in textwrap.wrap(paragraph, width=72):
                draw.text((80, y), line, fill="#334155", font=body_font)
                y += 38
                if y > 640:
                    break
            if y > 640:
                break
        draw.text((1135, 675), str(number), fill="#64748b", font=small_font)
        image.save(path, format="PNG")

    @staticmethod
    def _infer_title(lines: list[str], number: int) -> str:
        for line in lines:
            compact = clean_text(line)
            if 2 <= len(compact) <= 180:
                return compact
        return f"Slide {number}"

    @staticmethod
    def _slide_record(
        number: int,
        title: str,
        content: str,
        notes: str,
        images: list[str],
        visual_text: list[str],
        preview_image: str = "",
    ) -> dict[str, Any]:
        return {
            "slide": number,
            "title": title or f"Slide {number}",
            "content": content,
            "notes": notes,
            "images": images,
            "visual_text": visual_text,
            "preview_image": preview_image,
        }
