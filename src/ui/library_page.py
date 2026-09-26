"""Subject-library page."""

from __future__ import annotations

import base64
import binascii
from dataclasses import dataclass
import hashlib
import html
import json
from pathlib import Path
import re
import tempfile
from typing import Any

import streamlit as st

from ..jobs import JobError, JobManager, JobRecord
from ..library import LibraryDocument, LibraryError, LibraryFolder, LibraryStore
from .audio_transcript_component import render_audio_transcript
from .common import format_size, render_search_results, subject_lookup
from .file_manager_component import file_icon, render_file_manager


TEXT_PREVIEW_SUFFIXES = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".yaml",
    ".yml",
    ".py",
    ".html",
    ".htm",
    ".rtf",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
AUDIO_SUFFIXES = {".m4a", ".mp3", ".wav", ".aac", ".flac", ".ogg", ".opus"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".mkv", ".webm", ".m4v"}
# The explorer reports its 430px CSS minimum plus a 4px iframe allowance.
# Keep the adjacent preview on that same baseline.
INLINE_PREVIEW_HEIGHT = 434
INLINE_PREVIEW_CONTENT_HEIGHT = 390
DIRECT_DROP_MAX_FILE_BYTES = 64 * 1024 * 1024
DIRECT_DROP_MAX_TOTAL_BYTES = 96 * 1024 * 1024


@dataclass(frozen=True)
class _LectureSlideBundle:
    slides: list[dict[str, Any]]
    summaries: dict[int, str]
    alignment: dict[int, list[dict[str, Any]]]
    source_job: JobRecord | None


@dataclass(frozen=True)
class _LectureAudioBundle:
    paragraphs: list[dict[str, Any]]
    source_job: JobRecord | None


_INTERNAL_LECTURE_ARTIFACTS = (
    "_transcript_raw",
    "_transcript_cleaned",
    "_slides_extracted",
    "_slide_summaries",
    "_alignment",
    "_quality_report",
    "_manifest",
)


def _is_internal_lecture_artifact_name(filename: str) -> bool:
    """Recognize generated evidence that belongs behind a lecture preview."""
    path = Path(filename)
    normalized = re.sub(r"[^a-z0-9]+", "_", path.stem.casefold()).strip("_")
    return (
        "transcript" in normalized and path.suffix.lower() in {".json", ".txt"}
    ) or any(normalized.endswith(token.strip("_")) for token in _INTERNAL_LECTURE_ARTIFACTS) or (
        normalized.endswith("_notes") and path.suffix.lower() in {".md", ".markdown"}
    ) or ("_deep_review_" in f"_{normalized}_" and path.suffix.lower() in {".md", ".markdown"})


def _is_internal_lecture_artifact(document: LibraryDocument) -> bool:
    """Hide processing evidence from the learner-facing file explorer."""
    return _is_internal_lecture_artifact_name(document.original_name)


def _visible_documents(documents: list[LibraryDocument]) -> list[LibraryDocument]:
    return [document for document in documents if not _is_internal_lecture_artifact(document)]


def _load_json_file(path: Path | None) -> dict[str, Any]:
    if path is None or not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except (OSError, UnicodeError, json.JSONDecodeError):
        return {}


def _related_documents(library: LibraryStore, document: LibraryDocument) -> list[LibraryDocument]:
    return [
        item
        for item in library.list_documents(document.subject_id)
        if item.folder_id == document.folder_id
    ]


def _find_related_artifact(
    documents: list[LibraryDocument],
    required_tokens: tuple[str, ...],
) -> Path | None:
    for item in documents:
        normalized = re.sub(r"[^a-z0-9]+", "_", item.original_name.casefold()).strip("_")
        if all(token in normalized for token in required_tokens):
            return item.stored_path
    return None


def _source_names_match(stored_name: str, library_name: str) -> bool:
    stored = stored_name.casefold()
    target = library_name.casefold()
    return stored == target or any(
        stored.removeprefix(prefix) == target
        for prefix in ("slides_", "recording_")
    )


def _matches_document_content(path: Path, document: LibraryDocument) -> bool:
    try:
        if path.suffix.lower() != document.stored_path.suffix.lower() or path.stat().st_size != document.size_bytes:
            return False
        digest = hashlib.sha256()
        with path.open("rb") as source:
            for block in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(block)
        return digest.hexdigest() == document.sha256
    except OSError:
        return False


def _matching_lecture_job(manager: JobManager, document: LibraryDocument) -> JobRecord | None:
    """Resolve a library deck back to the completed lecture that created it."""
    name = document.original_name.casefold()
    folder = document.folder_name.casefold()
    for job in manager.list_jobs(limit=500):
        if job.kind != "lecture" or job.status != "completed":
            continue
        lecture_name = str(job.result.get("lecture_name", job.payload.get("lecture_name", ""))).casefold()
        lecture_title = str(job.result.get("lecture_title", job.payload.get("lecture_title", ""))).casefold()
        if lecture_name and name.startswith(lecture_name):
            return job
        if folder and lecture_title and folder == lecture_title:
            return job
        run_dir = Path(str(job.result.get("run_dir", "")))
        if run_dir.is_dir():
            try:
                input_files = list((run_dir / "input").iterdir())
                if any(_source_names_match(path.name, name) for path in input_files):
                    return job
                if any(_matches_document_content(path, document) for path in input_files):
                    return job
            except OSError:
                continue
    return None


def _result_path(job: JobRecord | None, key: str) -> Path | None:
    if job is None:
        return None
    value = str(job.result.get(key, "")).strip()
    return Path(value) if value else None


def _is_slide_deck(
    document: LibraryDocument,
    slides: list[dict[str, Any]],
    source_job: JobRecord | None,
) -> bool:
    suffix = document.stored_path.suffix.lower()
    if suffix in {".ppt", ".pptx"}:
        return True
    if suffix != ".pdf" or not slides or "notes" in document.original_name.casefold():
        return False
    normalized = re.sub(r"[^a-z0-9]+", "_", document.original_name.casefold())
    if "slides" in normalized:
        return True
    if source_job is not None:
        run_dir = Path(str(source_job.result.get("run_dir", "")))
        try:
            if any(
                _source_names_match(path.name, document.original_name)
                for path in (run_dir / "input").iterdir()
            ):
                return True
        except OSError:
            pass
    try:
        import fitz

        with fitz.open(document.stored_path) as pdf:
            return len(pdf) == len(slides)
    except Exception:
        return False


def _lecture_slide_bundle(
    library: LibraryStore,
    manager: JobManager,
    document: LibraryDocument,
) -> _LectureSlideBundle | None:
    if document.stored_path.suffix.lower() not in {".pdf", ".ppt", ".pptx"}:
        return None
    related = _related_documents(library, document)
    source_job = _matching_lecture_job(manager, document)
    slides_path = _find_related_artifact(related, ("slides", "extracted")) or _result_path(
        source_job, "slides_path"
    )
    slide_payload = _load_json_file(slides_path)
    slides = [item for item in slide_payload.get("slides", []) if isinstance(item, dict)]
    if not _is_slide_deck(document, slides, source_job):
        return None

    summaries_path = _find_related_artifact(related, ("slide", "summaries")) or _result_path(
        source_job, "slide_summaries_path"
    )
    summaries_payload = _load_json_file(summaries_path)
    summaries = {
        int(item["slide"]): str(item.get("summary", "")).strip()
        for item in summaries_payload.get("slides", [])
        if isinstance(item, dict) and str(item.get("slide", "")).isdigit()
    }
    alignment_path = _find_related_artifact(related, ("alignment",)) or _result_path(
        source_job, "alignment_path"
    )
    alignment_payload = _load_json_file(alignment_path)
    alignment = {
        int(item["slide"]): [value for value in item.get("paragraphs", []) if isinstance(value, dict)]
        for item in alignment_payload.get("slides", [])
        if isinstance(item, dict) and str(item.get("slide", "")).isdigit()
    }
    return _LectureSlideBundle(slides, summaries, alignment, source_job)


def _lecture_audio_bundle(
    library: LibraryStore,
    manager: JobManager,
    document: LibraryDocument,
) -> _LectureAudioBundle | None:
    if document.stored_path.suffix.lower() not in AUDIO_SUFFIXES:
        return None
    related = _related_documents(library, document)
    source_job = _matching_lecture_job(manager, document)
    transcript_path = next(
        (
            item.stored_path
            for item in related
            if "transcript_cleaned" in re.sub(
                r"[^a-z0-9]+", "_", item.original_name.casefold()
            ).strip("_")
            and item.stored_path.suffix.lower() == ".json"
        ),
        None,
    ) or _result_path(source_job, "transcript_path")
    payload = _load_json_file(transcript_path)
    metadata = payload.get("metadata", {}) if isinstance(payload.get("metadata"), dict) else {}
    transcript_kind = str(metadata.get("transcript_kind", "")).casefold()
    cleanup_enabled = True
    if source_job is not None and isinstance(source_job.payload.get("config"), dict):
        cleanup_enabled = bool(source_job.payload["config"].get("enable_transcript_cleanup", True))
    if transcript_kind and transcript_kind != "cleaned":
        return _LectureAudioBundle([], source_job)
    if source_job is not None and not cleanup_enabled:
        return _LectureAudioBundle([], source_job)
    paragraphs = [
        item
        for item in payload.get("paragraphs", [])
        if isinstance(item, dict) and str(item.get("text", "")).strip()
    ]
    return _LectureAudioBundle(paragraphs, source_job)


def _render_pdf(document: LibraryDocument, *, height: int = 680) -> None:
    """Render a stored PDF without exposing its local filesystem path to the browser."""
    try:
        pdf_data = document.stored_path.read_bytes()
    except OSError as exc:
        st.error(f"Could not open this PDF: {exc}")
        return

    try:
        st.pdf(pdf_data, height=height, key=f"pdf_viewer_{document.id}_{height}")
    except Exception:
        st.error(
            "The local PDF viewer is unavailable. Reinstall the project dependencies "
            "with pip install -r requirements.txt, then restart the website."
        )


def _read_text_preview(path: Path, maximum_bytes: int = 750_000) -> tuple[str, bool]:
    with path.open("rb") as source:
        data = source.read(maximum_bytes + 1)
    truncated = len(data) > maximum_bytes
    return data[:maximum_bytes].decode("utf-8", errors="replace"), truncated


@st.cache_data(show_spinner=False)
def _powerpoint_slide_images(path: str, modified_ns: int) -> dict[int, bytes]:
    """Cache rendered slide images without retaining temporary conversion files."""
    from ..pdf_processor import PDFProcessor

    with tempfile.TemporaryDirectory(prefix="library-ppt-preview-") as directory:
        image_dir = Path(directory) / "slide_images"
        image_dir.mkdir()
        images = PDFProcessor._render_powerpoint_previews(Path(path), image_dir)
        return {number: image.read_bytes() for number, image in images.items()}


def _render_powerpoint_preview(document: LibraryDocument) -> None:
    try:
        from pptx import Presentation

        deck = Presentation(document.stored_path)
        if not deck.slides:
            st.info("This presentation does not contain slides.")
            return
        with st.spinner("Preparing slide previews…"):
            images = _powerpoint_slide_images(
                str(document.stored_path), document.stored_path.stat().st_mtime_ns
            )
        if not images:
            st.caption("Slide images require LibreOffice. Showing extracted slide content below.")
        for slide_number, slide in enumerate(deck.slides, start=1):
            parts = [
                str(shape.text).strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and str(shape.text).strip()
            ]
            title = parts[0].splitlines()[0] if parts else f"Slide {slide_number}"
            with st.expander(f"Slide {slide_number}: {title}", expanded=slide_number == 1):
                if slide_number in images:
                    st.image(images[slide_number], width="stretch")
                st.markdown("#### Slide content")
                st.text("\n\n".join(parts) if parts else "(No extractable text on this slide.)")
    except Exception as exc:
        st.warning(f"Could not preview this PowerPoint file: {exc}")


def _limit_words(text: str, maximum: int) -> str:
    words = text.split()
    if len(words) <= maximum:
        return text.strip()
    return " ".join(words[:maximum]).rstrip(" ,;:") + "…"


def _exact_slide_summary(bundle: _LectureSlideBundle, slide: dict[str, Any]) -> str:
    number = int(slide.get("slide", 0))
    stored = bundle.summaries.get(number, "").strip()
    if stored:
        return _limit_words(stored, 90)

    slide_text = " ".join(str(slide.get("content", "")).split())
    professor_text = " ".join(
        " ".join(str(item.get("text", "")).split())
        for item in bundle.alignment.get(number, [])
        if str(item.get("text", "")).strip()
    )
    parts: list[str] = []
    if slide_text:
        parts.append(_limit_words(slide_text, 60))
    if professor_text and professor_text.casefold() not in slide_text.casefold():
        parts.append("Professor explanation: " + _limit_words(professor_text, 30))
    return " ".join(parts) or "No source-grounded summary is available for this slide."


def _render_current_slide(document: LibraryDocument, slide: dict[str, Any]) -> None:
    number = int(slide.get("slide", 1))
    suffix = document.stored_path.suffix.lower()
    if suffix == ".pdf":
        try:
            import fitz

            with fitz.open(document.stored_path) as pdf:
                page = pdf[number - 1]
                image = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0), alpha=False).tobytes("png")
            st.image(image, width="stretch")
            return
        except Exception as exc:
            st.warning(f"Could not render slide {number}: {exc}")

    preview_path = Path(str(slide.get("preview_image", "")))
    if preview_path.is_file():
        st.image(str(preview_path), width="stretch")
        return
    title = html.escape(str(slide.get("title", f"Slide {number}")))
    content = html.escape(str(slide.get("content", "")) or "(No extractable slide text.)")
    st.markdown(
        f'<div class="lecture-slide-fallback"><strong>{title}</strong><br><br>{content}</div>',
        unsafe_allow_html=True,
    )


def _matching_deep_reviews(
    manager: JobManager,
    source_job: JobRecord,
    slide_number: int,
) -> list[JobRecord]:
    matches = [
        job
        for job in manager.list_jobs(limit=500)
        if job.kind == "slide_review"
        and job.payload.get("source_job_id") == source_job.id
        and int(job.payload.get("slide_number", -1)) == slide_number
    ]
    return sorted(matches, key=lambda job: job.created_at, reverse=True)


def _render_deep_review_control(
    manager: JobManager,
    document: LibraryDocument,
    bundle: _LectureSlideBundle,
    slide_number: int,
    button_target: Any | None = None,
) -> None:
    source_job = bundle.source_job
    reviews = _matching_deep_reviews(manager, source_job, slide_number) if source_job else []
    latest = reviews[0] if reviews else None
    active = bool(latest and latest.status in {"queued", "running", "waiting", "deferred"})
    if latest and latest.status == "completed":
        markdown_path = Path(str(latest.result.get("markdown_path", "")))
        if markdown_path.is_file():
            with st.expander("Detailed explanation", expanded=True):
                st.markdown(markdown_path.read_text(encoding="utf-8"))
    elif latest and active:
        st.info(f"🧠 {latest.status.title()} · {latest.progress}% — {latest.message}")
        if st.button(
            "↻",
            help="Refresh review status",
            key=f"refresh_deep_review_{document.id}_{slide_number}",
            width="stretch",
        ):
            st.rerun()

    help_text = (
        "Generate a detailed, source-grounded explanation for only this slide."
        if source_job
        else "The completed lecture job for this file could not be located."
    )
    target = button_target or st
    if target.button(
        "🧠",
        type="primary",
        disabled=source_job is None or active,
        help=help_text,
        key=f"library_deep_review_{document.id}_{slide_number}",
        width="stretch",
    ):
        try:
            review_job = manager.enqueue_slide_review(
                source_job.id,
                slide_number,
                library_subject_id=document.subject_id,
                library_folder_id=document.folder_id,
                lecture_name=str(
                    source_job.result.get("lecture_name", source_job.payload.get("lecture_name", "Lecture"))
                ),
            )
            st.session_state["library_file_notice"] = (
                f"Deep Review for slide {slide_number} was queued as job {review_job.id[:8]}."
            )
            st.rerun()
        except JobError as exc:
            st.error(str(exc))


def _render_slide_document(
    manager: JobManager,
    document: LibraryDocument,
    bundle: _LectureSlideBundle,
    *,
    full_page: bool = False,
) -> None:
    slide_lookup = {int(item["slide"]): item for item in bundle.slides}
    slide_numbers = sorted(slide_lookup)
    slide_key = f"current_library_slide_{document.id}"
    try:
        selected_number = int(st.session_state.get(slide_key, slide_numbers[0]))
    except (TypeError, ValueError):
        selected_number = slide_numbers[0]
    if selected_number not in slide_lookup:
        selected_number = slide_numbers[0]
    st.session_state[slide_key] = selected_number
    slide = slide_lookup[selected_number]
    selected_index = slide_numbers.index(selected_number)
    st.html(
        """<style>
        .lecture-slide-nav { text-align:center; line-height:1.25; padding:.15rem .3rem; }
        .lecture-slide-nav strong { display:block; font-size:clamp(1rem, 1.3vw, 1.25rem); }
        .lecture-slide-nav span { color:#687083; font-size:.78rem; }
        .lecture-slide-summary { font-size:clamp(1.08rem, 1.45vw, 1.38rem); line-height:1.58; }
        .lecture-slide-fallback { font-size:clamp(1.08rem, 1.65vw, 1.55rem); line-height:1.55;
          padding:1.25rem; border:1px solid rgba(128,128,128,.25); border-radius:.75rem; }
        /* Streamlit's fit-content image toolbar can collapse in Safari. Give
           the slide an explicit CSS width instead of relying on its measured width. */
        .st-key-library_slide_image [data-testid="stFullScreenFrame"] > div:has(> [data-testid="stImage"]) {
          width: 100%;
        }
        .st-key-library_slide_image [data-testid="stImage"],
        .st-key-library_slide_image [data-testid="stImageContainer"] {
          width: 100%;
        }
        .st-key-library_slide_image [data-testid="stImageContainer"] > img {
          width: 100% !important; height: auto;
        }
        </style>""",
    )
    def render_slide() -> None:
        navigation = st.container(key="library_slide_navigation")
        previous_column, title_column, next_column = navigation.columns(
            [1, 5, 1], vertical_alignment="center"
        )
        if previous_column.button(
            "◀",
            disabled=selected_index == 0,
            help="Previous slide",
            key=f"previous_library_slide_{document.id}",
            width="stretch",
        ):
            st.session_state[slide_key] = slide_numbers[selected_index - 1]
            st.rerun()
        slide_title = html.escape(str(slide.get("title", "")).strip())
        title_column.markdown(
            f'<div class="lecture-slide-nav"><strong>{slide_title or f"Slide {selected_number}"}</strong>'
            f'<span>{selected_number} / {len(slide_numbers)}</span></div>',
            unsafe_allow_html=True,
        )
        if next_column.button(
            "▶",
            disabled=selected_index == len(slide_numbers) - 1,
            help="Next slide",
            key=f"next_library_slide_{document.id}",
            width="stretch",
        ):
            st.session_state[slide_key] = slide_numbers[selected_index + 1]
            st.rerun()
        with st.container(key="library_slide_image"):
            _render_current_slide(document, slide)

    def render_summary() -> None:
        summary_title, review_button = st.columns([5, 1], vertical_alignment="center")
        summary_title.markdown("#### Slide summary")
        summary = html.escape(_exact_slide_summary(bundle, slide))
        st.markdown(f'<div class="lecture-slide-summary">{summary}</div>', unsafe_allow_html=True)
        _render_deep_review_control(
            manager,
            document,
            bundle,
            selected_number,
            button_target=review_button,
        )

    if full_page:
        slide_column, summary_column = st.columns([1.65, 1], gap="large")
        with slide_column:
            render_slide()
        with summary_column:
            render_summary()
    else:
        render_slide()
        render_summary()

def _render_transcript_fallback(
    document: LibraryDocument,
    paragraphs: list[dict[str, Any]],
) -> None:
    st.audio(str(document.stored_path), format=document.media_type)
    if not paragraphs:
        st.info("A cleaned transcript is not available for this recording yet.")
        return
    timeline = st.container(height=560, border=True)
    with timeline:
        st.markdown("#### Cleaned transcript timeline")
        for paragraph in paragraphs:
            start = str(paragraph.get("start_time", "")).strip()
            end = str(paragraph.get("end_time", "")).strip()
            stamp = f"{start}–{end}".strip("–")
            if stamp:
                st.caption(stamp)
            st.write(str(paragraph.get("text", "")).strip())


def _render_audio_document(
    document: LibraryDocument,
    bundle: _LectureAudioBundle | None,
    *,
    full_page: bool = False,
) -> None:
    paragraphs = bundle.paragraphs if bundle is not None else []
    if not paragraphs:
        _render_transcript_fallback(document, paragraphs)
        return
    rendered = render_audio_transcript(
        document.stored_path,
        document.media_type,
        paragraphs,
        key=document.id,
        height=790 if full_page else INLINE_PREVIEW_CONTENT_HEIGHT,
    )
    if not rendered:
        _render_transcript_fallback(document, paragraphs)


def _render_document_content(
    document: LibraryDocument,
    audio_bundle: _LectureAudioBundle | None = None,
    *,
    full_page: bool = False,
) -> None:
    suffix = document.stored_path.suffix.lower()
    try:
        if suffix == ".pdf":
            _render_pdf(document, height=900 if full_page else INLINE_PREVIEW_CONTENT_HEIGHT)
        elif suffix in AUDIO_SUFFIXES:
            _render_audio_document(document, audio_bundle, full_page=full_page)
        elif suffix in VIDEO_SUFFIXES:
            st.video(str(document.stored_path), format=document.media_type)
        elif suffix in IMAGE_SUFFIXES:
            st.image(str(document.stored_path), width="stretch")
        elif suffix == ".json":
            text, truncated = _read_text_preview(document.stored_path)
            try:
                st.json(json.loads(text), expanded=1)
            except json.JSONDecodeError:
                st.code(text, language="json")
            if truncated:
                st.caption("Preview limited to the first 750 KB.")
        elif suffix in TEXT_PREVIEW_SUFFIXES:
            text, truncated = _read_text_preview(document.stored_path)
            language = suffix.lstrip(".")
            if language in {"md", "markdown"}:
                language = "markdown"
            st.code(text, language=language or None)
            if truncated:
                st.caption("Preview limited to the first 750 KB.")
        elif suffix == ".pptx":
            _render_powerpoint_preview(document)
        elif suffix == ".ppt":
            st.info("Legacy PPT files cannot be previewed directly. Download or convert this file to PPTX/PDF.")
        else:
            st.info("This file type has no inline preview. You can still download or organize the file.")
    except OSError as exc:
        st.error(f"Could not open this file: {exc}")


def _render_file_preview(
    library: LibraryStore,
    manager: JobManager,
    document: LibraryDocument,
    selected_key: str,
    slide_bundle: _LectureSlideBundle | None = None,
    audio_bundle: _LectureAudioBundle | None = None,
) -> None:
    # Style-only HTML is placed outside the content flow by Streamlit.
    st.html(
        """<style>
        .st-key-library_preview_content {
          position: relative;
        }
        /* Remove the layout wrapper, not just its keyed child, from the flow. */
        .st-key-library_preview_content > [data-testid="stLayoutWrapper"]:has(> .st-key-library_preview_menu) {
          position: absolute; top: .4rem; right: .4rem; z-index: 100;
          width: 2.7rem; min-width: 0; height: auto; overflow: visible;
        }
        .st-key-library_preview_menu {
          width: 100%; overflow: visible;
        }
        .st-key-library_preview_content .st-key-library_slide_navigation {
          padding-right: 2.7rem;
        }
        .st-key-library_preview_menu div[data-testid="stPopover"] > button {
          border-radius: 999px; background: color-mix(in srgb, var(--secondary-background-color) 88%, transparent);
          box-shadow: 0 2px 10px rgba(0,0,0,.22); font-size: 1.15rem;
        }
        </style>""",
    )
    with st.container(key="library_preview_menu"):
        with st.popover("⋯", help="File actions", width="stretch"):
            if st.button(
                "Close preview",
                key=f"close_preview_{document.id}",
                width="stretch",
            ):
                st.session_state.pop(selected_key, None)
                st.rerun()
            st.link_button(
                "Full screen ↗",
                f"?open_file={document.id}",
                help="Open this file in a full-page browser view.",
                width="stretch",
            )
            with st.expander("Rename"):
                with st.form(f"rename_file_{document.id}"):
                    new_name = st.text_input("File name", value=document.original_name)
                    rename = st.form_submit_button("Save", width="stretch")
                if rename:
                    try:
                        updated = library.rename_document(document.id, new_name)
                        st.session_state["library_file_notice"] = f"Renamed file to {updated.original_name}."
                        st.rerun()
                    except LibraryError as exc:
                        st.error(str(exc))
            if document.size_bytes <= 100 * 1024 * 1024:
                try:
                    st.download_button(
                        "Download",
                        data=document.stored_path.read_bytes(),
                        file_name=document.original_name,
                        mime=document.media_type,
                        key=f"download_file_{document.id}",
                        width="stretch",
                    )
                except OSError as exc:
                    st.error(f"Could not prepare this file for download: {exc}")
            else:
                st.caption("This file is too large for an in-browser download.")
            if st.button(
                "Delete",
                key=f"delete_preview_{document.id}",
                width="stretch",
            ):
                st.session_state["library_document_delete_confirmation"] = document.id

            if st.session_state.get("library_document_delete_confirmation") == document.id:
                st.warning(f"Permanently delete '{document.original_name}'?")
                confirm, cancel = st.columns(2)
                if confirm.button("Delete", type="primary", key=f"confirm_delete_{document.id}"):
                    try:
                        library.delete_document(document.id)
                        st.session_state.pop("library_document_delete_confirmation", None)
                        st.session_state.pop(selected_key, None)
                        st.session_state["library_file_notice"] = f"Deleted {document.original_name}."
                        st.rerun()
                    except LibraryError as exc:
                        st.error(str(exc))
                if cancel.button("Cancel", key=f"cancel_delete_{document.id}"):
                    st.session_state.pop("library_document_delete_confirmation", None)
                    st.rerun()

    if document.extraction_error:
        st.warning(document.extraction_error)
    if slide_bundle is not None:
        _render_slide_document(manager, document, slide_bundle)
    else:
        _render_document_content(document, audio_bundle)


def render_full_library_document(
    library: LibraryStore,
    manager: JobManager,
    document_id: str,
) -> None:
    """Render one library file without the workspace chrome for a new tab."""
    try:
        document = library.get_document(document_id)
    except LibraryError as exc:
        st.error(str(exc))
        return

    st.title(f"{file_icon(document.original_name)} {document.original_name}")
    st.caption(
        f"{document.subject_name} / {document.folder_name or 'Subject root'} · "
        f"{format_size(document.size_bytes)}"
    )
    st.divider()
    slide_bundle = _lecture_slide_bundle(library, manager, document)
    if slide_bundle is not None:
        _render_slide_document(manager, document, slide_bundle, full_page=True)
        return
    audio_bundle = _lecture_audio_bundle(library, manager, document)
    _render_document_content(document, audio_bundle, full_page=True)


def _handle_file_manager_event(
    library: LibraryStore,
    subject_id: str,
    folders: list[LibraryFolder],
    documents: list[LibraryDocument],
    browse_key: str,
    selected_key: str,
    event: dict[str, object] | None,
) -> None:
    if not event:
        return
    event_id = str(event.get("event_id", ""))
    handled_key = f"library_handled_explorer_event_{subject_id}"
    if not event_id or st.session_state.get(handled_key) == event_id:
        return
    st.session_state[handled_key] = event_id

    action = str(event.get("action", ""))
    folder_id = str(event.get("folder_id", ""))
    document_id = str(event.get("document_id", ""))
    folder_lookup = {folder.id: folder for folder in folders}
    document_lookup = {document.id: document for document in documents}

    if action == "browse":
        if folder_id and folder_id not in folder_lookup:
            st.error("That folder no longer exists. Refresh the library and try again.")
            return
        st.session_state[browse_key] = folder_id
        st.session_state.pop(selected_key, None)
        st.rerun()

    if action == "open":
        if document_id not in document_lookup:
            st.error("That file no longer exists. Refresh the library and try again.")
            return
        st.session_state[selected_key] = document_id
        st.rerun()

    if action == "create_folder":
        try:
            folder = library.create_folder(subject_id, str(event.get("name", "")))
            st.session_state[browse_key] = folder.id
            st.session_state["library_file_notice"] = f"Created {folder.name}."
            st.rerun()
        except LibraryError as exc:
            st.error(str(exc))
        return

    if action == "rename_folder":
        if folder_id not in folder_lookup:
            st.error("That folder no longer exists. Refresh the library and try again.")
            return
        try:
            renamed = library.rename_folder(folder_id, str(event.get("new_name", "")))
            st.session_state["library_file_notice"] = f"Renamed folder to {renamed.name}."
            st.rerun()
        except LibraryError as exc:
            st.error(str(exc))
        return

    if action == "delete_folder":
        if folder_id not in folder_lookup:
            st.error("That folder no longer exists. Refresh the library and try again.")
            return
        folder = folder_lookup[folder_id]
        try:
            library.delete_folder(folder.id, delete_documents=bool(folder.document_count))
            st.session_state[browse_key] = ""
            st.session_state.pop(selected_key, None)
            st.session_state["library_file_notice"] = f"Deleted {folder.name}."
            st.rerun()
        except LibraryError as exc:
            st.error(str(exc))
        return

    if action == "upload":
        if folder_id and folder_id not in folder_lookup:
            st.error("That folder no longer exists. Refresh the library and try again.")
            return
        raw_files = event.get("files", [])
        if not isinstance(raw_files, list) or not raw_files:
            st.error("No files were received. Try dropping the files again.")
            return

        uploaded = 0
        errors: list[str] = []
        skipped_too_large = event.get("skipped_too_large", [])
        if isinstance(skipped_too_large, list):
            errors.extend(
                f"{str(filename)} is too large to upload by direct drop."
                for filename in skipped_too_large
                if isinstance(filename, str) and filename.strip()
            )
        total_bytes = 0
        for item in raw_files:
            if not isinstance(item, dict):
                errors.append("One dropped item could not be read.")
                continue
            filename = str(item.get("name", "")).strip()
            encoded_data = item.get("data", "")
            if not filename or not isinstance(encoded_data, str):
                errors.append("One dropped file is missing its name or data.")
                continue
            try:
                data = base64.b64decode(encoded_data, validate=True)
            except (binascii.Error, ValueError):
                errors.append(f"Could not read {filename}.")
                continue
            if len(data) > DIRECT_DROP_MAX_FILE_BYTES:
                errors.append(f"{filename} is too large to upload by direct drop.")
                continue
            if total_bytes + len(data) > DIRECT_DROP_MAX_TOTAL_BYTES:
                errors.append("The dropped files are too large to upload together. Drop fewer files at once.")
                continue
            total_bytes += len(data)
            try:
                library.add_document_bytes(subject_id, filename, data, folder_id=folder_id or None)
                uploaded += 1
            except LibraryError as exc:
                errors.append(f"Could not add {filename}: {exc}")

        destination = folder_lookup[folder_id].name if folder_id else "Subject root"
        if uploaded:
            st.session_state["library_file_notice"] = (
                f"Added {uploaded} file{'s' if uploaded != 1 else ''} to {destination}."
            )
        if errors:
            st.session_state["library_file_upload_errors"] = errors
        st.rerun()

    if action != "move":
        return
    if document_id not in document_lookup or (folder_id and folder_id not in folder_lookup):
        st.error("That drag-and-drop move is no longer valid. Refresh the library and try again.")
        return

    target_name = folder_lookup[folder_id].name if folder_id else "Subject root"
    try:
        moved = library.move_document_to_folder(document_id, folder_id or None)
        st.session_state.pop(selected_key, None)
        st.session_state["library_file_notice"] = f"Moved {moved.original_name} to {target_name}."
        st.rerun()
    except LibraryError as exc:
        st.error(str(exc))


def _render_subject_selector(library: LibraryStore) -> str | None:
    creating_key = "library_creating_subject"
    name_key = "library_new_subject_name"
    error_key = "library_new_subject_error"
    notice_key = "library_new_subject_notice"

    if st.session_state.get(creating_key, False):
        def create_subject() -> None:
            name = str(st.session_state.get(name_key, "")).strip()
            if not name:
                st.session_state[error_key] = "Subject name cannot be empty."
                return
            try:
                subject = library.create_subject(name)
            except LibraryError as exc:
                st.session_state[error_key] = str(exc)
                return
            st.session_state["library_open_subject"] = subject.id
            st.session_state[creating_key] = False
            st.session_state.pop(error_key, None)
            st.session_state[notice_key] = f"Created {subject.name}."

        st.text_input(
            "Subject name",
            key=name_key,
            max_chars=120,
            placeholder="Type a subject name and press Enter",
            icon=":material/add:",
            on_change=create_subject,
        )
        error = st.session_state.get(error_key)
        if error:
            st.error(str(error))
        return None

    notice = st.session_state.pop(notice_key, None)
    if notice:
        st.success(str(notice))

    subjects = library.list_subjects()
    lookup = subject_lookup(subjects)
    selector, add = st.columns([12, 1], vertical_alignment="bottom")
    with selector:
        selected_id = st.selectbox(
            "Open subject",
            options=[subject.id for subject in subjects],
            format_func=lambda value: lookup[value].name,
            key="library_open_subject",
            placeholder="No subjects yet",
            disabled=not subjects,
        )
    with add:
        if st.button(
            "＋",
            key="library_add_subject",
            help="Add a subject",
            width="stretch",
        ):
            st.session_state[creating_key] = True
            st.session_state.pop(name_key, None)
            st.session_state.pop(error_key, None)
            st.rerun()
    if not subjects:
        st.info("Select + to add your first subject.")
        return None
    return str(selected_id)


def render_library(library: LibraryStore, manager: JobManager) -> None:
    st.title("📚 Subject Library")
    st.caption(
        "Browse folders like a desktop file manager, drag files to move them, "
        "and click a file to preview its contents."
    )

    selected_id = _render_subject_selector(library)
    if selected_id is None:
        return
    subjects = library.list_subjects()
    lookup = subject_lookup(subjects)
    subject = lookup[selected_id]
    st.subheader(subject.name)
    if subject.description:
        st.write(subject.description)

    folders = library.list_folders(subject.id)
    documents = _visible_documents(library.list_documents(subject.id))
    folder_lookup = {folder.id: folder for folder in folders}
    document_lookup = {document.id: document for document in documents}
    browse_key = f"library_browse_folder_{subject.id}"
    selected_key = f"library_selected_file_{subject.id}"
    if st.session_state.get(browse_key, "") not in {"", *folder_lookup}:
        st.session_state[browse_key] = ""
    if st.session_state.get(selected_key) not in document_lookup:
        st.session_state.pop(selected_key, None)
    current_folder_id = st.session_state.get(browse_key, "")
    selected_document_id = st.session_state.get(selected_key)
    query = st.text_input(
        "Search this subject",
        placeholder="Search this subject",
        icon=":material/search:",
        label_visibility="collapsed",
        key=f"library_search_{subject.id}",
    )

    notice = st.session_state.pop("library_file_notice", "")
    if notice:
        st.success(notice)
    upload_errors = st.session_state.pop("library_file_upload_errors", [])
    for error in upload_errors:
        st.error(error)

    if query.strip():
        results = [
            result
            for result in library.search_documents(query, subject_id=subject.id, limit=30)
            if not _is_internal_lecture_artifact_name(result.document_name)
        ][:12]
        render_search_results(results)

    selected_document = document_lookup.get(selected_document_id) if selected_document_id else None
    slide_bundle = (
        _lecture_slide_bundle(library, manager, selected_document)
        if selected_document is not None
        else None
    )
    audio_bundle = (
        _lecture_audio_bundle(library, manager, selected_document)
        if selected_document is not None
        else None
    )
    explorer_column, preview_column = st.columns([3, 2], gap="small")
    with explorer_column:
        manager_event = render_file_manager(
            folders,
            documents,
            current_folder_id=current_folder_id or None,
            selected_document_id=selected_document_id,
            key=f"library_explorer_{subject.id}",
        )
    with preview_column:
        with st.container(
            height=INLINE_PREVIEW_HEIGHT,
            border=True,
            key="library_preview_content",
        ):
            if selected_document_id and selected_document_id in document_lookup:
                _render_file_preview(
                    library,
                    manager,
                    document_lookup[selected_document_id],
                    selected_key,
                    slide_bundle,
                    audio_bundle,
                )
            else:
                st.info("Click a file icon to preview its contents here.")

    _handle_file_manager_event(
        library,
        subject.id,
        folders,
        documents,
        browse_key,
        selected_key,
        manager_event,
    )
