"""Persistent subject and document library for the local knowledge workspace."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import hashlib
import mimetypes
from pathlib import Path
import re
import shutil
import sqlite3
from typing import Any
from uuid import uuid4

from .embeddings import chunk_text
from .pdf_runtime import muted_mupdf_errors
from .utils import clean_text, safe_filename


class LibraryError(RuntimeError):
    """Raised when a library operation cannot be completed safely."""


class DuplicateDocumentError(LibraryError):
    """Raised when the same file already exists in the selected subject."""


@dataclass(frozen=True)
class Subject:
    id: str
    name: str
    description: str
    created_at: str


@dataclass(frozen=True)
class LibraryDocument:
    id: str
    subject_id: str
    subject_name: str
    original_name: str
    stored_path: Path
    media_type: str
    size_bytes: int
    sha256: str
    status: str
    extraction_error: str
    created_at: str


@dataclass(frozen=True)
class SearchResult:
    document_id: str
    document_name: str
    subject_id: str
    subject_name: str
    locator: str
    text: str
    score: float

    @property
    def citation(self) -> str:
        return f"{self.document_name} — {self.locator}"


class LibraryStore:
    """Own subject folders, document metadata, extracted text, and search."""

    TEXT_SUFFIXES = {
        ".txt",
        ".md",
        ".markdown",
        ".csv",
        ".tsv",
        ".json",
        ".yaml",
        ".yml",
        ".py",
        ".ipynb",
        ".html",
        ".htm",
    }
    SEARCH_STOP_WORDS = {
        "about",
        "and",
        "are",
        "can",
        "could",
        "document",
        "documents",
        "find",
        "for",
        "from",
        "have",
        "into",
        "please",
        "show",
        "that",
        "the",
        "this",
        "what",
        "where",
        "which",
        "with",
        "would",
    }

    def __init__(self, root: str | Path) -> None:
        self.root = Path(root).resolve()
        self.subjects_root = self.root / "subjects"
        self.database_path = self.root / "library.sqlite3"
        self.subjects_root.mkdir(parents=True, exist_ok=True)
        self._initialize_schema()

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.database_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute("PRAGMA journal_mode = WAL")
        return connection

    def _initialize_schema(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS subjects (
                    id TEXT PRIMARY KEY,
                    name TEXT NOT NULL COLLATE NOCASE UNIQUE,
                    description TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS documents (
                    id TEXT PRIMARY KEY,
                    subject_id TEXT NOT NULL REFERENCES subjects(id) ON DELETE CASCADE,
                    original_name TEXT NOT NULL,
                    stored_path TEXT NOT NULL UNIQUE,
                    media_type TEXT NOT NULL,
                    size_bytes INTEGER NOT NULL,
                    sha256 TEXT NOT NULL,
                    status TEXT NOT NULL,
                    extraction_error TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    UNIQUE(subject_id, sha256)
                );

                CREATE TABLE IF NOT EXISTS document_chunks (
                    id TEXT PRIMARY KEY,
                    document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE,
                    ordinal INTEGER NOT NULL,
                    locator TEXT NOT NULL,
                    content TEXT NOT NULL,
                    UNIQUE(document_id, ordinal)
                );

                CREATE INDEX IF NOT EXISTS idx_documents_subject ON documents(subject_id);
                CREATE INDEX IF NOT EXISTS idx_chunks_document ON document_chunks(document_id);
                """
            )

    def create_subject(self, name: str, description: str = "") -> Subject:
        normalized_name = clean_text(name)
        normalized_description = clean_text(description)
        if not normalized_name:
            raise LibraryError("Subject name cannot be empty.")
        if len(normalized_name) > 120:
            raise LibraryError("Subject name must contain no more than 120 characters.")
        subject_id = uuid4().hex
        created_at = self._timestamp()
        try:
            with self._connect() as connection:
                connection.execute(
                    "INSERT INTO subjects(id, name, description, created_at) VALUES (?, ?, ?, ?)",
                    (subject_id, normalized_name, normalized_description, created_at),
                )
                (self.subjects_root / subject_id / "files").mkdir(parents=True, exist_ok=False)
                (self.subjects_root / subject_id / "derived").mkdir(parents=True, exist_ok=False)
        except sqlite3.IntegrityError as exc:
            raise LibraryError(f"A subject named '{normalized_name}' already exists.") from exc
        except OSError as exc:
            raise LibraryError(f"Could not create the subject folder: {exc}") from exc
        return Subject(subject_id, normalized_name, normalized_description, created_at)

    def list_subjects(self) -> list[Subject]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT subjects.*, COUNT(documents.id) AS document_count
                FROM subjects
                LEFT JOIN documents ON documents.subject_id = subjects.id
                GROUP BY subjects.id
                ORDER BY subjects.name COLLATE NOCASE
                """
            ).fetchall()
        return [self._subject_from_row(row) for row in rows]

    def get_subject(self, subject_id: str) -> Subject:
        with self._connect() as connection:
            row = connection.execute("SELECT * FROM subjects WHERE id = ?", (subject_id,)).fetchone()
        if row is None:
            raise LibraryError("The selected subject no longer exists.")
        return self._subject_from_row(row)

    def add_document(
        self,
        subject_id: str,
        source_path: str | Path,
        filename: str | None = None,
    ) -> LibraryDocument:
        source = Path(source_path)
        if not source.is_file():
            raise LibraryError(f"Document file not found: {source}")
        subject = self.get_subject(subject_id)
        original_name = safe_filename(filename or source.name, "document")
        checksum_builder = hashlib.sha256()
        size_bytes = 0
        try:
            with source.open("rb") as input_file:
                while block := input_file.read(1024 * 1024):
                    checksum_builder.update(block)
                    size_bytes += len(block)
        except OSError as exc:
            raise LibraryError(f"Could not read '{original_name}': {exc}") from exc
        if size_bytes == 0:
            raise LibraryError(f"'{original_name}' is empty and was not added.")
        checksum = checksum_builder.hexdigest()
        self._require_not_duplicate(subject, checksum)

        document_id = uuid4().hex
        relative_path = Path("subjects") / subject_id / "files" / f"{document_id}_{original_name}"
        destination = self.root / relative_path
        temporary = destination.with_name(destination.name + ".part")
        try:
            shutil.copy2(source, temporary)
            temporary.replace(destination)
        except OSError as exc:
            if temporary.exists():
                temporary.unlink()
            raise LibraryError(f"Could not store '{original_name}': {exc}") from exc
        return self._index_and_register(
            subject,
            document_id,
            original_name,
            relative_path,
            destination,
            checksum,
            size_bytes,
        )

    def add_document_bytes(self, subject_id: str, filename: str, data: bytes) -> LibraryDocument:
        subject = self.get_subject(subject_id)
        original_name = safe_filename(filename, "document")
        if not data:
            raise LibraryError(f"'{original_name}' is empty and was not added.")
        checksum = hashlib.sha256(data).hexdigest()
        self._require_not_duplicate(subject, checksum)

        document_id = uuid4().hex
        relative_path = Path("subjects") / subject_id / "files" / f"{document_id}_{original_name}"
        destination = self.root / relative_path
        temporary = destination.with_name(destination.name + ".part")
        try:
            temporary.write_bytes(data)
            temporary.replace(destination)
        except OSError as exc:
            if temporary.exists():
                temporary.unlink()
            raise LibraryError(f"Could not store '{original_name}': {exc}") from exc

        return self._index_and_register(
            subject,
            document_id,
            original_name,
            relative_path,
            destination,
            checksum,
            len(data),
        )

    def _require_not_duplicate(self, subject: Subject, checksum: str) -> None:
        with self._connect() as connection:
            duplicate = connection.execute(
                "SELECT original_name FROM documents WHERE subject_id = ? AND sha256 = ?",
                (subject.id, checksum),
            ).fetchone()
        if duplicate:
            raise DuplicateDocumentError(
                f"This file is already stored in {subject.name} as '{duplicate['original_name']}'."
            )

    def _index_and_register(
        self,
        subject: Subject,
        document_id: str,
        original_name: str,
        relative_path: Path,
        destination: Path,
        checksum: str,
        size_bytes: int,
    ) -> LibraryDocument:
        """Extract searchable chunks and atomically register a stored file."""

        media_type = mimetypes.guess_type(original_name)[0] or "application/octet-stream"
        status = "stored"
        extraction_error = ""
        chunks: list[tuple[str, str]] = []
        try:
            sections = self._extract_sections(destination)
            chunks = self._chunk_sections(sections)
            if chunks:
                status = "indexed"
        except Exception as exc:
            status = "extraction_failed"
            extraction_error = str(exc)[:1000]

        created_at = self._timestamp()
        try:
            with self._connect() as connection:
                connection.execute(
                    """
                    INSERT INTO documents(
                        id, subject_id, original_name, stored_path, media_type,
                        size_bytes, sha256, status, extraction_error, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        document_id,
                        subject.id,
                        original_name,
                        str(relative_path),
                        media_type,
                        size_bytes,
                        checksum,
                        status,
                        extraction_error,
                        created_at,
                    ),
                )
                connection.executemany(
                    "INSERT INTO document_chunks(id, document_id, ordinal, locator, content) VALUES (?, ?, ?, ?, ?)",
                    [
                        (f"{document_id}:{ordinal}", document_id, ordinal, locator, content)
                        for ordinal, (locator, content) in enumerate(chunks, start=1)
                    ],
                )
        except sqlite3.IntegrityError as exc:
            if destination.exists():
                destination.unlink()
            raise DuplicateDocumentError(f"'{original_name}' is already registered in this subject.") from exc
        except sqlite3.Error as exc:
            if destination.exists():
                destination.unlink()
            raise LibraryError(f"Could not register '{original_name}' in the library: {exc}") from exc
        return LibraryDocument(
            document_id,
            subject.id,
            subject.name,
            original_name,
            destination,
            media_type,
            size_bytes,
            checksum,
            status,
            extraction_error,
            created_at,
        )

    def list_documents(self, subject_id: str | None = None) -> list[LibraryDocument]:
        parameters: tuple[Any, ...] = ()
        where = ""
        if subject_id:
            where = "WHERE documents.subject_id = ?"
            parameters = (subject_id,)
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT documents.*, subjects.name AS subject_name
                FROM documents
                JOIN subjects ON subjects.id = documents.subject_id
                {where}
                ORDER BY documents.created_at DESC
                """,
                parameters,
            ).fetchall()
        return [self._document_from_row(row) for row in rows]

    def get_document(self, document_id: str) -> LibraryDocument:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT documents.*, subjects.name AS subject_name
                FROM documents
                JOIN subjects ON subjects.id = documents.subject_id
                WHERE documents.id = ?
                """,
                (document_id,),
            ).fetchone()
        if row is None:
            raise LibraryError("The selected document no longer exists.")
        return self._document_from_row(row)

    def rename_document(self, document_id: str, new_name: str) -> LibraryDocument:
        """Rename one stored file without invalidating its indexed content."""
        document = self.get_document(document_id)
        normalized = safe_filename(new_name, "")
        if not normalized:
            raise LibraryError("The new file name cannot be empty.")
        current_suffix = Path(document.original_name).suffix
        requested_suffix = Path(normalized).suffix
        if current_suffix and not requested_suffix:
            normalized += current_suffix
            requested_suffix = current_suffix
        if current_suffix.lower() != requested_suffix.lower():
            raise LibraryError(
                f"Keep the original '{current_suffix or '(no extension)'}' file extension when renaming this document."
            )
        if normalized == document.original_name:
            return document
        with self._connect() as connection:
            duplicate = connection.execute(
                """
                SELECT 1 FROM documents
                WHERE subject_id = ? AND id != ? AND original_name = ? COLLATE NOCASE
                """,
                (document.subject_id, document.id, normalized),
            ).fetchone()
        if duplicate:
            raise LibraryError(f"A document named '{normalized}' already exists in {document.subject_name}.")

        relative_path = Path("subjects") / document.subject_id / "files" / f"{document.id}_{normalized}"
        destination = self.root / relative_path
        try:
            document.stored_path.replace(destination)
            with self._connect() as connection:
                connection.execute(
                    "UPDATE documents SET original_name = ?, stored_path = ?, media_type = ? WHERE id = ?",
                    (
                        normalized,
                        str(relative_path),
                        mimetypes.guess_type(normalized)[0] or "application/octet-stream",
                        document.id,
                    ),
                )
        except (OSError, sqlite3.Error) as exc:
            if destination.exists() and not document.stored_path.exists():
                destination.replace(document.stored_path)
            raise LibraryError(f"Could not rename '{document.original_name}': {exc}") from exc
        return self.get_document(document.id)

    def move_document(self, document_id: str, target_subject_id: str) -> LibraryDocument:
        """Move a document and its search ownership to another subject."""
        document = self.get_document(document_id)
        target = self.get_subject(target_subject_id)
        if target.id == document.subject_id:
            return document
        with self._connect() as connection:
            duplicate_content = connection.execute(
                "SELECT original_name FROM documents WHERE subject_id = ? AND sha256 = ?",
                (target.id, document.sha256),
            ).fetchone()
            duplicate_name = connection.execute(
                "SELECT 1 FROM documents WHERE subject_id = ? AND original_name = ? COLLATE NOCASE",
                (target.id, document.original_name),
            ).fetchone()
        if duplicate_content:
            raise DuplicateDocumentError(
                f"This file already exists in {target.name} as '{duplicate_content['original_name']}'."
            )
        if duplicate_name:
            raise LibraryError(f"A document named '{document.original_name}' already exists in {target.name}.")

        relative_path = Path("subjects") / target.id / "files" / f"{document.id}_{document.original_name}"
        destination = self.root / relative_path
        try:
            document.stored_path.replace(destination)
            with self._connect() as connection:
                connection.execute(
                    "UPDATE documents SET subject_id = ?, stored_path = ? WHERE id = ?",
                    (target.id, str(relative_path), document.id),
                )
        except (OSError, sqlite3.Error) as exc:
            if destination.exists() and not document.stored_path.exists():
                destination.replace(document.stored_path)
            raise LibraryError(f"Could not move '{document.original_name}' to {target.name}: {exc}") from exc
        return self.get_document(document.id)

    def search_documents(
        self,
        query: str,
        subject_id: str | None = None,
        limit: int = 10,
    ) -> list[SearchResult]:
        raw_terms = [item.lower() for item in re.findall(r"[\w'-]{2,}", clean_text(query))]
        terms = [item for item in raw_terms if item not in self.SEARCH_STOP_WORDS] or raw_terms
        if not terms:
            return []
        parameters: list[Any] = []
        where = ""
        if subject_id:
            where = "WHERE documents.subject_id = ?"
            parameters.append(subject_id)
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT COALESCE(document_chunks.content, '') AS content,
                       COALESCE(document_chunks.locator, 'Stored file') AS locator,
                       documents.id AS document_id, documents.original_name,
                       subjects.id AS subject_id, subjects.name AS subject_name
                FROM documents
                JOIN subjects ON subjects.id = documents.subject_id
                LEFT JOIN document_chunks ON documents.id = document_chunks.document_id
                {where}
                """,
                tuple(parameters),
            ).fetchall()

        phrase = clean_text(query).lower()
        ranked: list[SearchResult] = []
        for row in rows:
            content = str(row["content"])
            lowered = content.lower()
            name = str(row["original_name"]).lower()
            subject_name = str(row["subject_name"]).lower()
            content_hits = sum(lowered.count(term) for term in terms)
            metadata_hits = sum((2 if term in name else 0) + (1 if term in subject_name else 0) for term in terms)
            phrase_bonus = 6 if len(phrase) > 3 and phrase in lowered else 0
            matched_terms = sum(term in lowered or term in name for term in set(terms))
            score = float(content_hits + metadata_hits + phrase_bonus + matched_terms * 0.5)
            if score <= 0:
                continue
            ranked.append(
                SearchResult(
                    document_id=str(row["document_id"]),
                    document_name=str(row["original_name"]),
                    subject_id=str(row["subject_id"]),
                    subject_name=str(row["subject_name"]),
                    locator=str(row["locator"]),
                    text=(
                        self._search_snippet(content, terms)
                        if content
                        else "(The file is stored, but it does not have searchable extracted text yet.)"
                    ),
                    score=score,
                )
            )
        ranked.sort(key=lambda item: (-item.score, item.document_name.lower(), item.locator))
        return ranked[: max(1, min(limit, 50))]

    def stats(self) -> dict[str, int]:
        with self._connect() as connection:
            subjects = connection.execute("SELECT COUNT(*) FROM subjects").fetchone()[0]
            documents = connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
            indexed = connection.execute("SELECT COUNT(*) FROM documents WHERE status = 'indexed'").fetchone()[0]
        return {"subjects": int(subjects), "documents": int(documents), "indexed_documents": int(indexed)}

    def _extract_sections(self, path: Path) -> list[tuple[str, str]]:
        suffix = path.suffix.lower()
        if suffix in self.TEXT_SUFFIXES:
            text = path.read_text(encoding="utf-8", errors="replace")
            return [("Text", text)] if clean_text(text) else []
        if suffix == ".pdf":
            try:
                import fitz
            except ImportError as exc:
                raise LibraryError("PyMuPDF is required to index PDF documents.") from exc
            sections: list[tuple[str, str]] = []
            try:
                with muted_mupdf_errors(fitz), fitz.open(path) as document:
                    for page_number, page in enumerate(document, start=1):
                        text = page.get_text("text")
                        if clean_text(text):
                            sections.append((f"Page {page_number}", text))
            except Exception as exc:
                raise LibraryError(f"Could not extract PDF text: {exc}") from exc
            return sections
        if suffix == ".pptx":
            try:
                from pptx import Presentation
            except ImportError as exc:
                raise LibraryError("python-pptx is required to index PowerPoint documents.") from exc
            try:
                deck = Presentation(path)
                sections = []
                for slide_number, slide in enumerate(deck.slides, start=1):
                    parts = [
                        clean_text(shape.text)
                        for shape in slide.shapes
                        if getattr(shape, "has_text_frame", False) and clean_text(shape.text)
                    ]
                    if parts:
                        sections.append((f"Slide {slide_number}", "\n".join(parts)))
                return sections
            except Exception as exc:
                raise LibraryError(f"Could not extract PowerPoint text: {exc}") from exc
        return []

    @staticmethod
    def _chunk_sections(sections: list[tuple[str, str]]) -> list[tuple[str, str]]:
        output: list[tuple[str, str]] = []
        for locator, text in sections:
            pieces = chunk_text(text, size=1_200, overlap=160)
            for index, piece in enumerate(pieces, start=1):
                part = f" · part {index}" if len(pieces) > 1 else ""
                output.append((f"{locator}{part}", piece))
        return output

    @staticmethod
    def _search_snippet(content: str, terms: list[str], maximum: int = 700) -> str:
        if len(content) <= maximum:
            return content
        lowered = content.lower()
        positions = [lowered.find(term) for term in terms if lowered.find(term) >= 0]
        center = min(positions) if positions else 0
        start = max(0, center - maximum // 4)
        end = min(len(content), start + maximum)
        prefix = "…" if start else ""
        suffix = "…" if end < len(content) else ""
        return prefix + content[start:end].strip() + suffix

    @staticmethod
    def _timestamp() -> str:
        return datetime.now(timezone.utc).isoformat()

    @staticmethod
    def _subject_from_row(row: sqlite3.Row) -> Subject:
        return Subject(str(row["id"]), str(row["name"]), str(row["description"]), str(row["created_at"]))

    def _document_from_row(self, row: sqlite3.Row) -> LibraryDocument:
        return LibraryDocument(
            id=str(row["id"]),
            subject_id=str(row["subject_id"]),
            subject_name=str(row["subject_name"]),
            original_name=str(row["original_name"]),
            stored_path=self.root / str(row["stored_path"]),
            media_type=str(row["media_type"]),
            size_bytes=int(row["size_bytes"]),
            sha256=str(row["sha256"]),
            status=str(row["status"]),
            extraction_error=str(row["extraction_error"]),
            created_at=str(row["created_at"]),
        )
